from pathlib import Path
import json,csv,hashlib,io
import openpyxl,pdfplumber
from pypdf import PdfReader,PdfWriter
from pypdf.annotations import Rectangle
from pypdf.generic import ArrayObject,FloatObject,NameObject
R=Path(__file__).resolve().parents[1]
x=R/'raw_data/PG_2026_Company_Financials.xlsx';p=R/'raw_data/PG_2026_Annual_Report.pdf'
s=openpyxl.load_workbook(x,data_only=True).worksheets[0]
rows=[]
for i in list(range(5,18))+[20,21]:
 vals=[s.cell(i,c).value for c in [6,4,2]]
 vals=[0 if isinstance(v,str) and v.strip()=='—' else v for v in vals]
 name=s.cell(i,1).value
 if i in [20,21]:name+=' EPS'
 rows.append({'metric':name,'values':vals,'type':'Reported','unit':'USD/share' if i in [20,21] else 'USD million','company_excel_row':i})
with pdfplumber.open(p) as pdf:
 page=pdf.pages[48];text=page.extract_text()
 for row in rows:
  for v in row['values']:
   if v and abs(v)>=10:assert f'{abs(v):,}' in text,(row['metric'],v)
 (R/'raw_data/earnings_page.txt').write_text(text)
 words=page.extract_words()
reader=PdfReader(p);out=PdfWriter();out.add_page(reader.pages[48]);out.write(R/'raw_data/PG_2026_Earnings_Original_Page.pdf')
marked=PdfWriter();marked.add_page(reader.pages[48]);height=float(reader.pages[48].mediabox.height)
for w in words:
 if w['text'] in ['87,032','84,284','84,039','19,748','20,451','18,545']:
  a=Rectangle(rect=(w['x0']-2,height-w['bottom']-2,w['x1']+2,height-w['top']+2))
  a[NameObject('/C')]=ArrayObject([FloatObject(.85),FloatObject(.05),FloatObject(.05)])
  a[NameObject('/Border')]=ArrayObject([FloatObject(0),FloatObject(0),FloatObject(1.3)])
  marked.add_annotation(page_number=0,annotation=a)
marked.write(R/'evidence/source_annotated.pdf')
(R/'processed_data/financial_history.json').write_text(json.dumps(rows,indent=2)+'\n')
with (R/'processed_data/financial_history.csv').open('w') as f:
 w=csv.writer(f);w.writerow(['metric','fiscal_year','value','unit','classification','source_page','source_excel_row'])
 for row in rows:
  for year,v in zip([2024,2025,2026],row['values']):w.writerow([row['metric'],year,v,row['unit'],'Reported',37,row['company_excel_row']])
manifest={'sources':[{'file':str(q.relative_to(R)),'sha256':hashlib.sha256(q.read_bytes()).hexdigest()} for q in [p,x]],'report_url':'https://s204.q4cdn.com/332108499/files/doc_financials/2026/ar/2026_annual_report.pdf','company_excel_url':'https://s204.q4cdn.com/332108499/files/doc_financials/2026/ar/PG-2026-Financial-Statements-for-IR-site.xlsx','printed_page':37,'pdf_page':49,'periods':[2024,2025,2026],'fiscal_year_end':'June 30','reported_values':45,'normalisation':'Source em dashes on impairment become numeric zero; negatives retained; years reordered chronologically. No adjusted earnings calculated.'}
(R/'raw_data/source_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print('45 reported inputs extracted; numeric PDF cross-check passed')
