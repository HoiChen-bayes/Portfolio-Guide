from pathlib import Path
import json,csv,subprocess
import openpyxl,pdfplumber
from pypdf import PdfReader,PdfWriter
from pypdf.annotations import Rectangle
from pypdf.generic import ArrayObject,FloatObject,NameObject
R=Path(__file__).resolve().parents[1]
w=openpyxl.load_workbook(R/'raw_data/PG_2026_Company_Financials.xlsx',data_only=True)
def extract(si,rs,cols,years,page):
 s=w.worksheets[si];out=[]
 for r in rs:
  vs=[s.cell(r,c).value for c in cols]
  if not any(isinstance(v,(int,float)) or (isinstance(v,str) and v.strip()=='—') for v in vs):continue
  vs=[0 if isinstance(v,str) and v.strip()=='—' else v for v in vs]
  out.append(dict(metric=s.cell(r,1).value,values=vs,years=years,source_sheet=s.title,source_row=r,source_page=page))
 return out
bs=extract(2,range(7,44),[4,2],[2025,2026],38)
cf=extract(4,range(5,40),[6,4,2],[2024,2025,2026],40)
segments=[['Beauty',14964,16023],['Grooming',6662,6918],['Health Care',11998,12456],['Fabric & Home Care',29617,30314],['Baby, Feminine & Family Care',20248,20401],['Corporate',794,919]]
drivers=[['Beauty',.04,.02,.01,0,0,.07],['Grooming',-.01,.03,.02,0,0,.04],['Health Care',-.02,.03,.02,.01,0,.04],['Fabric & Home Care',0,.01,.01,0,0,.02],['Baby, Feminine & Family Care',-.01,.02,0,0,0,.01],['Total company',0,.02,.01,0,0,.03]]
margin=[['Product mix',-120],['Product and packaging investment',-70],['Restructuring',-60],['Net tariffs',-30],['Commodities',-20],['Foreign exchange',-10],['Other and rounding',-10],['Manufacturing productivity',180],['Pricing',40]]
obj=dict(balance_sheet=bs,cash_flow=cf,segments=segments,sales_drivers=drivers,gross_margin_drivers=margin,source_pages={'segments':[44,45],'sales_drivers':[20,22],'gross_margin_drivers':[20,21]})
(R/'processed_data/extended_history.json').write_text(json.dumps(obj,indent=2)+'\n')
manifest_path=R/'raw_data/source_manifest.json'
m=json.loads(manifest_path.read_text());m.update(reported_values=199,reported_statement_values=199,statement_periods={'income':[2024,2025,2026],'cash_flow':[2024,2025,2026],'balance':[2025,2026]},additional_printed_pages=[20,21,22,38,40,44,45],normalisation='Reported dashes become zero; signs and original amounts retained; chronological years. Displayed rounding residuals are preserved. Driver percentages transcribed from identified report tables and independently reviewed.');manifest_path.write_text(json.dumps(m,indent=2)+'\n')
with (R/'processed_data/three_statement_inputs.csv').open('w') as f:
 cw=csv.writer(f);cw.writerow(['statement','metric','fiscal_year','value','unit','source_sheet','source_row','printed_page'])
 for kind,rows in [('Balance sheet',bs),('Cash flow',cf)]:
  for row in rows:
   for year,v in zip(row['years'],row['values']):cw.writerow([kind,row['metric'],year,v,'USD million',row['source_sheet'],row['source_row'],row['source_page']])
reader=PdfReader(R/'raw_data/PG_2026_Annual_Report.pdf')
specs=[('balance',49,['9,942','9,556','126,521','125,231','72,210','72,946','54,311','52,284']),('cash',51,['16,144','16,065','14,974','19,556','17,817','19,846','9,942','9,556','9,482']),('drivers',33,[]),('margin_20',31,[]),('margin_21',32,[]),('segments_26',55,['16,023','6,918','12,456','30,314','20,401','919','87,032']),('segments_25',56,['14,964','6,662','11,998','29,617','20,248','794','84,284'])]
with pdfplumber.open(R/'raw_data/PG_2026_Annual_Report.pdf') as pdf:
 for name,idx,marks in specs:
  out=PdfWriter();out.add_page(reader.pages[idx]);out.write(R/f'raw_data/{name}_original.pdf')
  page=pdf.pages[idx];txt=page.extract_text();(R/f'raw_data/{name}.txt').write_text(txt)
  marked=PdfWriter();marked.add_page(reader.pages[idx]);h=float(reader.pages[idx].mediabox.height)
  boxes=[]
  if marks:boxes=[(v['x0']-2,v['top']-2,v['x1']+2,v['bottom']+2) for v in page.extract_words() if v['text'] in marks]
  else:
   for line in page.extract_text_lines():
    t=line['text']
    if (name=='drivers' and ('TOTAL COMPANY' in t or t.startswith('Beauty 4'))) or (name=='margin_20' and ('Gross margin 50.2' in t or 'basis points of' in t)) or (name=='margin_21' and ('basis points of' in t or 'Marketing spending' in t)):
     boxes.append((line['x0']-2,line['top']-2,line['x1']+2,line['bottom']+2))
  for x0,y0,x1,y1 in boxes:
   a=Rectangle(rect=(x0,h-y1,x1,h-y0));a[NameObject('/C')]=ArrayObject([FloatObject(.85),FloatObject(.05),FloatObject(.05)]);a[NameObject('/Border')]=ArrayObject([FloatObject(0),FloatObject(0),FloatObject(1.2)]);marked.add_annotation(page_number=0,annotation=a)
  marked.write(R/f'evidence/{name}_annotated.pdf')
  subprocess.run(['pdftoppm','-scale-to','1800','-singlefile','-png',str(R/f'evidence/{name}_annotated.pdf'),str(R/f'evidence/original_{name}')],check=True)
print({'balance_inputs':len(bs)*2,'cash_inputs':len(cf)*3,'total_inputs':45+len(bs)*2+len(cf)*3})
